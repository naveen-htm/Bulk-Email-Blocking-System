import csv
import io
import os
from datetime import datetime
from functools import wraps

from flask import Flask, flash, jsonify, redirect, render_template, request, session, url_for
from flask_sqlalchemy import SQLAlchemy
from werkzeug.security import check_password_hash, generate_password_hash

app = Flask(__name__)
app.config["SECRET_KEY"] = os.getenv("SECRET_KEY", "change-this-secret-key")
app.config["SQLALCHEMY_DATABASE_URI"] = os.getenv("DATABASE_URL", "sqlite:///bulk_email.db")
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
db = SQLAlchemy(app)


class User(db.Model):
    __tablename__ = "users"
    user_id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(50), unique=True, nullable=False)
    password_hash = db.Column(db.String(255), nullable=False)
    role = db.Column(db.String(20), nullable=False, default="ADMIN")
    created_at = db.Column(db.DateTime, default=datetime.utcnow)


class EmailAccount(db.Model):
    __tablename__ = "email_accounts"
    email_id = db.Column(db.Integer, primary_key=True)
    email_address = db.Column(db.String(255), unique=True, nullable=False)
    domain = db.Column(db.String(255), nullable=False)
    account_type = db.Column(db.String(50), nullable=False, default="STUDENT")
    status = db.Column(db.String(20), nullable=False, default="ACTIVE")
    created_at = db.Column(db.DateTime, default=datetime.utcnow)


class BlockLog(db.Model):
    __tablename__ = "block_logs"
    log_id = db.Column(db.Integer, primary_key=True)
    email_id = db.Column(db.Integer, db.ForeignKey("email_accounts.email_id"), nullable=False)
    action = db.Column(db.String(20), nullable=False)
    performed_by = db.Column(db.Integer, db.ForeignKey("users.user_id"), nullable=False)
    action_time = db.Column(db.DateTime, default=datetime.utcnow)
    email = db.relationship("EmailAccount")
    user = db.relationship("User")


class OperationLog(db.Model):
    __tablename__ = "operation_logs"
    operation_id = db.Column(db.Integer, primary_key=True)
    operation_type = db.Column(db.String(20), nullable=False)
    total_accounts = db.Column(db.Integer, default=0)
    successful_accounts = db.Column(db.Integer, default=0)
    failed_accounts = db.Column(db.Integer, default=0)
    performed_by = db.Column(db.Integer, db.ForeignKey("users.user_id"), nullable=False)
    operation_time = db.Column(db.DateTime, default=datetime.utcnow)
    user = db.relationship("User")


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if "user_id" not in session:
            return redirect(url_for("login"))
        return view(*args, **kwargs)
    return wrapped


def valid_email(value):
    return isinstance(value, str) and "@" in value and "." in value.rsplit("@", 1)[-1]


def normalize_email(value):
    return str(value).strip().lower()


def add_operation(op_type, total, success, failed):
    log = OperationLog(operation_type=op_type, total_accounts=total,
                       successful_accounts=success, failed_accounts=failed,
                       performed_by=session["user_id"])
    db.session.add(log)
    db.session.commit()
    return log


@app.route("/")
def index():
    return redirect(url_for("dashboard") if "user_id" in session else url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        user = User.query.filter_by(username=username).first()
        if user and check_password_hash(user.password_hash, password):
            session["user_id"] = user.user_id
            session["username"] = user.username
            session["role"] = user.role
            return redirect(url_for("dashboard"))
        flash("Invalid username or password.", "danger")
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/dashboard")
@login_required
def dashboard():
    total = EmailAccount.query.count()
    active = EmailAccount.query.filter_by(status="ACTIVE").count()
    blocked = EmailAccount.query.filter_by(status="BLOCKED").count()
    operations = OperationLog.query.order_by(OperationLog.operation_time.desc()).limit(8).all()
    return render_template("dashboard.html", total=total, active=active, blocked=blocked, operations=operations)


@app.route("/emails")
@login_required
def emails():
    search = request.args.get("search", "").strip()
    domain = request.args.get("domain", "").strip()
    status = request.args.get("status", "").strip()
    account_type = request.args.get("account_type", "").strip()
    query = EmailAccount.query
    if search:
        query = query.filter(EmailAccount.email_address.ilike(f"%{search}%"))
    if domain:
        query = query.filter_by(domain=domain)
    if status:
        query = query.filter_by(status=status)
    if account_type:
        query = query.filter_by(account_type=account_type)
    accounts = query.order_by(EmailAccount.email_id.desc()).all()
    domains = [x[0] for x in db.session.query(EmailAccount.domain).distinct().order_by(EmailAccount.domain)]
    types = [x[0] for x in db.session.query(EmailAccount.account_type).distinct().order_by(EmailAccount.account_type)]
    return render_template("emails.html", accounts=accounts, domains=domains, types=types,
                           search=search, selected_domain=domain, selected_status=status,
                           selected_type=account_type)


@app.route("/upload", methods=["POST"])
@login_required
def upload():
    file = request.files.get("file")
    if not file or not file.filename:
        flash("Please select a CSV or XLSX file.", "danger")
        return redirect(url_for("emails"))
    filename = file.filename.lower()
    rows = []
    try:
        if filename.endswith(".csv"):
            text = file.stream.read().decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text))
            if not reader.fieldnames:
                raise ValueError("The CSV file has no header row.")
            rows = list(reader)
            columns = {str(c).strip().lower(): c for c in reader.fieldnames if c}
        elif filename.endswith(".xlsx"):
            from openpyxl import load_workbook
            workbook = load_workbook(file, read_only=True, data_only=True)
            sheet = workbook.active
            values = list(sheet.iter_rows(values_only=True))
            if not values:
                raise ValueError("The Excel file is empty.")
            headers = [str(value).strip() if value is not None else "" for value in values[0]]
            columns = {header.lower(): header for header in headers if header}
            rows = [dict(zip(headers, row)) for row in values[1:]]
        else:
            flash("Only CSV and XLSX files are supported.", "danger")
            return redirect(url_for("emails"))
    except Exception as exc:
        flash(f"Could not read the file: {exc}", "danger")
        return redirect(url_for("emails"))

    email_col = columns.get("email") or columns.get("email_address") or columns.get("email address")
    type_col = columns.get("account_type") or columns.get("type") or columns.get("account type")
    if not email_col:
        flash("File must contain an 'email' or 'email_address' column.", "danger")
        return redirect(url_for("emails"))

    added = duplicates = invalid = 0
    for row in rows:
        raw_email = row.get(email_col, "")
        email = normalize_email(raw_email if raw_email is not None else "")
        if not valid_email(email):
            invalid += 1
            continue
        if EmailAccount.query.filter_by(email_address=email).first():
            duplicates += 1
            continue
        account_type = "STUDENT"
        raw_type = row.get(type_col) if type_col else None
        if raw_type is not None and str(raw_type).strip():
            account_type = str(raw_type).strip().upper()[:50]
        domain = email.rsplit("@", 1)[1]
        db.session.add(EmailAccount(email_address=email, domain=domain,
                                    account_type=account_type, status="ACTIVE"))
        added += 1
    db.session.commit()
    flash(f"Upload complete: {added} added, {duplicates} duplicates, {invalid} invalid.", "success")
    return redirect(url_for("emails"))


def perform_action(action, ids):
    accounts = EmailAccount.query.filter(EmailAccount.email_id.in_(ids)).all()
    total = len(ids)
    success = failed = 0
    target = "BLOCKED" if action == "BLOCK" else "ACTIVE"
    for account in accounts:
        if account.status == target:
            failed += 1
            continue
        account.status = target
        db.session.add(BlockLog(email_id=account.email_id, action=action,
                                 performed_by=session["user_id"]))
        success += 1
    failed += max(0, total - len(accounts))
    db.session.commit()
    add_operation(action, total, success, failed)
    return total, success, failed


@app.route("/bulk-action", methods=["POST"])
@login_required
def bulk_action():
    action = request.form.get("action", "").upper()
    ids = [int(x) for x in request.form.getlist("email_ids") if x.isdigit()]
    if action not in ("BLOCK", "UNBLOCK"):
        flash("Invalid operation.", "danger")
        return redirect(url_for("emails"))
    if not ids:
        flash("Select at least one email account.", "warning")
        return redirect(url_for("emails"))
    total, success, failed = perform_action(action, ids)
    flash(f"{action.title()} complete: {success} successful, {failed} failed, {total} selected.", "success")
    return redirect(url_for("emails"))


@app.route("/domain-action", methods=["POST"])
@login_required
def domain_action():
    action = request.form.get("action", "").upper()
    domain = request.form.get("domain", "").strip().lower()
    if action not in ("BLOCK", "UNBLOCK") or not domain:
        flash("Invalid domain operation.", "danger")
        return redirect(url_for("emails"))
    accounts = EmailAccount.query.filter_by(domain=domain).all()
    ids = [x.email_id for x in accounts]
    if not ids:
        flash("No accounts found for that domain.", "warning")
        return redirect(url_for("emails"))
    total, success, failed = perform_action(action, ids)
    flash(f"Domain {domain}: {action.title()} complete — {success} successful, {failed} failed.", "success")
    return redirect(url_for("emails"))


@app.route("/logs")
@login_required
def logs():
    logs = BlockLog.query.order_by(BlockLog.action_time.desc()).all()
    return render_template("logs.html", logs=logs)


@app.route("/reports")
@login_required
def reports():
    operations = OperationLog.query.order_by(OperationLog.operation_time.desc()).all()
    block_count = BlockLog.query.filter_by(action="BLOCK").count()
    unblock_count = BlockLog.query.filter_by(action="UNBLOCK").count()
    return render_template("reports.html", operations=operations,
                           block_count=block_count, unblock_count=unblock_count)


@app.route("/api/stats")
@login_required
def stats():
    return jsonify({"total": EmailAccount.query.count(),
                    "active": EmailAccount.query.filter_by(status="ACTIVE").count(),
                    "blocked": EmailAccount.query.filter_by(status="BLOCKED").count()})


@app.cli.command("init-db")
def init_db():
    db.create_all()
    if not User.query.filter_by(username="admin").first():
        db.session.add(User(username="admin", password_hash=generate_password_hash("admin123"), role="ADMIN"))
        db.session.commit()
    print("Database initialized. Login: admin / admin123")


@app.cli.command("seed")
def seed():
    db.create_all()
    if not User.query.filter_by(username="admin").first():
        db.session.add(User(username="admin", password_hash=generate_password_hash("admin123"), role="ADMIN"))
        db.session.commit()
    samples = [("student1@gmail.com", "STUDENT"), ("student2@gmail.com", "STUDENT"),
               ("student3@gmail.com", "STUDENT"), ("staff1@yahoo.com", "STAFF"),
               ("staff2@yahoo.com", "STAFF"), ("admin1@college.com", "ADMIN"),
               ("student4@college.com", "STUDENT"), ("student5@college.com", "STUDENT")]
    for email, typ in samples:
        if not EmailAccount.query.filter_by(email_address=email).first():
            db.session.add(EmailAccount(email_address=email, domain=email.rsplit("@", 1)[1], account_type=typ))
    db.session.commit()
    print("Sample data inserted.")


if __name__ == "__main__":
    with app.app_context():
        db.create_all()
        if not User.query.filter_by(username="admin").first():
            db.session.add(User(username="admin", password_hash=generate_password_hash("admin123"), role="ADMIN"))
            db.session.commit()
    app.run(debug=True)
